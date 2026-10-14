"""Shared fixtures. External network access is blocked so every test proves it
runs offline. Loopback (127.0.0.1 / ::1) is allowed because the Prefect test
harness talks to a temporary local server."""
from __future__ import annotations

import socket
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

FIXTURES = Path(__file__).parent / "fixtures"
LOOPBACK_HOSTS = {"127.0.0.1", "::1", "localhost"}
_real_connect = socket.socket.connect


def _is_loopback(address) -> bool:
    if isinstance(address, (str, bytes)):  # AF_UNIX path
        return True
    return isinstance(address, tuple) and address[0] in LOOPBACK_HOSTS


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    def guard(self, address):
        if _is_loopback(address):
            return _real_connect(self, address)
        raise RuntimeError(f"Network access attempted during tests: {address!r}")

    monkeypatch.setattr(socket.socket, "connect", guard)


@pytest.fixture
def sdmx_csv_bytes() -> bytes:
    return (FIXTURES / "sdmx_cpi_sample.csv").read_bytes()


def build_abs_workbook(path: Path) -> Path:
    """A tiny workbook in the ABS 'Time Series Workbook' layout (2 series, 3 quarters)."""
    wb = openpyxl.Workbook()
    wb.active.title = "Index"
    ws = wb.create_sheet("Data1")
    ws.append([None, "Index Numbers ;  Rents ;  Sydney ;", "Index Numbers ;  Rents ;  Darwin ;"])
    ws.append(["Unit", "Index Numbers", "Index Numbers"])
    ws.append(["Series Type", "Original", "Original"])
    ws.append(["Data Type", "INDEX", "INDEX"])
    ws.append(["Frequency", "Quarter", "Quarter"])
    ws.append(["Collection Month", 3, 3])
    ws.append(["Series Start", datetime(1980, 6, 1), datetime(1980, 9, 1)])
    ws.append(["Series End", datetime(1980, 12, 1), datetime(1980, 12, 1)])
    ws.append(["No. Obs", 3, 2])
    ws.append(["Series ID", "A0000001A", "A0000002B"])
    ws.append([datetime(1980, 6, 1), 10.5, None])  # Darwin not yet published
    ws.append([datetime(1980, 9, 1), 10.7, 9.9])
    ws.append([datetime(1980, 12, 1), 10.9, 10.1])
    ws.append([None, None, None])  # trailing blank row, as in real files
    wb.create_sheet("Enquiries")
    wb.save(path)
    return path


@pytest.fixture
def abs_workbook(tmp_path) -> Path:
    return build_abs_workbook(tmp_path / "sample.xlsx")
