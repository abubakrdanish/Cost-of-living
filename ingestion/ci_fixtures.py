"""Build and check the small raw-data fixtures that CI runs `dbt build` against.

CI never calls the ABS. Instead it loads tests/fixtures/ci_raw/, which is cut
from real downloads in data/raw/:

* CSV fixtures keep the real header and a subset of the real lines, byte for
  byte (no values are edited), from 2020-Q3 onward, plus a few rows the
  staging models must filter out (other index codes, seasonally adjusted,
  other WPI sectors).
* The archive XLSX fixture keeps the real ABS workbook layout and metadata
  rows for the Rents series and one unrelated series, from 2020-Q3 onward.
* One deliberately edited, older rents snapshot is added so CI exercises the
  revision and row-count checks. It is the only edited file; see
  tests/fixtures/ci_raw/README.md.

Keeping fixtures in sync with the real schema:
    python -m ingestion.ci_fixtures check      # compare fixture layout with data/raw
    python -m ingestion.ci_fixtures generate   # rebuild fixtures from data/raw
The manual "live ingestion" GitHub workflow runs `check` after a real
download, so a change in the ABS file layout shows up there.
"""
from __future__ import annotations

import argparse
import csv
import io
import sys
from collections.abc import Callable
from datetime import datetime
from pathlib import Path

import openpyxl

from ingestion import cpi, rents, wpi
from ingestion.abs_timeseries_xlsx import FIRST_DATA_ROW, METADATA_ROWS
from ingestion.config import PROJECT_ROOT, RAW_DATA_DIR
from ingestion.dataset import RawDataset
from ingestion.raw_files import list_data_files

FIXTURE_DIR = PROJECT_ROOT / "tests" / "fixtures" / "ci_raw"
FIRST_QUARTER = "2020-Q3"
FIRST_MONTH = "2020-07"
FIRST_ARCHIVE_PERIOD = datetime(2020, 9, 1)  # ABS stamps 2020-Q3 as Sep 2020

# The older, edited rents snapshot used to exercise revision handling.
OLDER_RENTS_FILE = "abs_cpi_rents_20260915.csv"
REVISED_KEY = ("30014", "10", "2", "M", "2026-07")  # Rents, Original, Melbourne, Jul 2026
REVISED_OFFSET = -0.5

Row = dict[str, str]


def _code(value: str) -> str:
    """'2: Melbourne' -> '2'."""
    return value.split(":", 1)[0].strip()


def _row(header: list[str], line: str) -> Row:
    values = next(csv.reader(io.StringIO(line)))
    return {h.split(":", 1)[0].strip(): v for h, v in zip(header, values)}


def _keep_cpi(row: Row) -> bool:
    if _code(row["TSEST"]) != "10":
        return False
    if _code(row["FREQ"]) == "M":
        return row["TIME_PERIOD"] >= FIRST_MONTH
    return row["TIME_PERIOD"] >= FIRST_QUARTER


def _keep_rents(row: Row) -> bool:
    headline = _code(row["INDEX"]) == "30014" and _code(row["TSEST"]) == "10"
    if not headline:
        # Decoys staging must drop (code 115522, seasonally adjusted): 2026 only.
        return row["TIME_PERIOD"] >= "2026"
    if _code(row["FREQ"]) == "M":
        return row["TIME_PERIOD"] >= FIRST_MONTH
    return row["TIME_PERIOD"] >= FIRST_QUARTER


def _keep_wpi(row: Row) -> bool:
    headline = (_code(row["INDEX"]) == "THRPEB" and _code(row["SECTOR"]) == "7"
                and _code(row["INDUSTRY"]) == "TOT" and _code(row["TSEST"]) == "10")
    if headline:
        return row["TIME_PERIOD"] >= FIRST_QUARTER
    # A few decoys staging must filter out: Victoria private sector, Australia SA.
    decoy = ((_code(row["SECTOR"]) == "1" and _code(row["REGION"]) == "2"
              and _code(row["INDUSTRY"]) == "TOT" and _code(row["TSEST"]) == "10")
             or (_code(row["INDEX"]) == "THRPEB" and _code(row["SECTOR"]) == "7"
                 and _code(row["INDUSTRY"]) == "TOT" and _code(row["TSEST"]) == "20"))
    return decoy and row["TIME_PERIOD"] >= "2025-Q3"


CSV_FILTERS: dict[str, Callable[[Row], bool]] = {
    cpi.CPI_ALL_GROUPS.name: _keep_cpi,
    rents.CPI_RENTS.name: _keep_rents,
    wpi.WPI_INDEX.name: _keep_wpi,
}
ARCHIVE = rents.CPI_CAPITAL_CITIES_QUARTERLY_ARCHIVE
ARCHIVE_EXTRA_DESCRIPTION = "Index Numbers ;  All groups CPI ;  Melbourne ;"


def all_datasets() -> list[RawDataset]:
    return cpi.DATASETS + rents.DATASETS + wpi.DATASETS


def latest_file(dataset: RawDataset, root: Path) -> Path:
    files = list_data_files(dataset.raw_dir(root), dataset.name, dataset.ext)
    if not files:
        raise FileNotFoundError(f"No raw file for {dataset.name} under {root}; run ingestion first")
    return files[-1]


# ---- CSV --------------------------------------------------------------------

def trim_csv(text: str, keep: Callable[[Row], bool]) -> str:
    """Header plus the original lines (line endings included) whose row passes `keep`."""
    lines = text.splitlines(keepends=True)
    header = next(csv.reader(io.StringIO(lines[0])))
    kept = [line for line in lines[1:] if line.strip() and keep(_row(header, line))]
    return lines[0] + "".join(kept)


def make_older_rents_snapshot(text: str) -> str:
    """An earlier download: without the newest month, and one value different.

    Every other line is copied unchanged. The edited line is rebuilt from its
    fields, which is safe because Rents lines contain no quoted fields.
    """
    lines = text.splitlines(keepends=True)
    header = next(csv.reader(io.StringIO(lines[0])))
    idx = {name.split(":", 1)[0]: i for i, name in enumerate(header)}
    parsed = [(line, next(csv.reader(io.StringIO(line)))) for line in lines[1:] if line.strip()]
    newest_month = max(f[idx["TIME_PERIOD"]] for _, f in parsed if _code(f[idx["FREQ"]]) == "M")
    out, edited = [lines[0]], 0
    for line, f in parsed:
        key = (_code(f[idx["INDEX"]]), _code(f[idx["TSEST"]]), _code(f[idx["REGION"]]),
               _code(f[idx["FREQ"]]), f[idx["TIME_PERIOD"]])
        if key[3] == "M" and key[4] == newest_month:
            continue
        if key == REVISED_KEY:
            if '"' in line:
                raise ValueError("Unexpected quoted field in the line to edit")
            ending = line[len(line.rstrip("\r\n")):]
            f[idx["OBS_VALUE"]] = f"{float(f[idx['OBS_VALUE']]) + REVISED_OFFSET:.2f}"
            line = ",".join(f) + ending
            edited += 1
        out.append(line)
    if edited != 1:
        raise ValueError(f"Expected to edit exactly 1 row for {REVISED_KEY}, edited {edited}")
    return "".join(out)


# ---- XLSX -------------------------------------------------------------------

def _selected_columns(description_row: tuple) -> list[int]:
    return [0] + [i for i, d in enumerate(description_row)
                  if d and (";  Rents ;" in d or d.strip() == ARCHIVE_EXTRA_DESCRIPTION.strip())]


def trim_archive_workbook(source: Path, target: Path) -> None:
    """Copy the ABS layout keeping Rents (+1 decoy) series from 2020-Q3."""
    src = openpyxl.load_workbook(source, read_only=True, data_only=True)
    out = openpyxl.Workbook()
    out.active.title = "Index"
    out["Index"]["A1"] = "CI fixture cut from 6401.0 Sep quarter 2025 Table 9 (see README.md)"
    sheet_no = 0
    for name in src.sheetnames:
        if not name.startswith("Data"):
            continue
        rows = list(src[name].iter_rows(values_only=True))
        cols = _selected_columns(rows[0])
        if len(cols) == 1:
            continue
        sheet_no += 1
        ws = out.create_sheet(f"Data{sheet_no}")
        for row in rows[:FIRST_DATA_ROW]:
            ws.append([row[c] for c in cols])
        for row in rows[FIRST_DATA_ROW:]:
            if isinstance(row[0], datetime) and row[0] >= FIRST_ARCHIVE_PERIOD:
                ws.append([row[c] for c in cols])
    src.close()
    out.save(target)


# ---- Commands -----------------------------------------------------------------

def generate(raw_root: Path = RAW_DATA_DIR, fixture_root: Path = FIXTURE_DIR) -> list[Path]:
    """Rebuild every fixture from the newest real raw files."""
    written = []
    for dataset in all_datasets():
        source = latest_file(dataset, raw_root)
        target_dir = dataset.raw_dir(fixture_root)
        target_dir.mkdir(parents=True, exist_ok=True)
        target = target_dir / source.name
        if dataset.kind == "sdmx_csv":
            # Bytes in, bytes out: keeps ABS's CRLF line endings exactly.
            trimmed = trim_csv(source.read_bytes().decode("utf-8"), CSV_FILTERS[dataset.name])
            target.write_bytes(trimmed.encode("utf-8"))
            if dataset.name == rents.CPI_RENTS.name:
                older = target_dir / OLDER_RENTS_FILE
                older.write_bytes(make_older_rents_snapshot(trimmed).encode("utf-8"))
                written.append(older)
        else:
            trim_archive_workbook(source, target)
        written.append(target)
    return written


def csv_header(path: Path) -> list[str]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return next(csv.reader(handle))


def xlsx_layout(path: Path) -> list[str]:
    """Metadata row labels of the first Data sheet (what the parser depends on)."""
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        sheet = next(wb[n] for n in wb.sheetnames if n.startswith("Data"))
        rows = sheet.iter_rows(min_row=2, max_row=FIRST_DATA_ROW, max_col=1, values_only=True)
        return [r[0] for r in rows]
    finally:
        wb.close()


def check(raw_root: Path = RAW_DATA_DIR, fixture_root: Path = FIXTURE_DIR) -> list[str]:
    """Return a list of differences between real raw files and the fixtures."""
    problems = []
    for dataset in all_datasets():
        try:
            real = latest_file(dataset, raw_root)
            fixture = latest_file(dataset, fixture_root)
        except FileNotFoundError as exc:
            problems.append(str(exc))
            continue
        if dataset.kind == "sdmx_csv":
            real_cols, fix_cols = csv_header(real), csv_header(fixture)
        else:
            real_cols, fix_cols = xlsx_layout(real), xlsx_layout(fixture)
            if real_cols != list(METADATA_ROWS):
                problems.append(f"{dataset.name}: real file metadata rows changed: {real_cols}")
        if real_cols != fix_cols:
            problems.append(f"{dataset.name}: layout differs\n  real:    {real_cols}\n  fixture: {fix_cols}")
    return problems


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("command", choices=["generate", "check"])
    args = parser.parse_args(argv)
    if args.command == "generate":
        for path in generate():
            print(f"wrote {path.relative_to(PROJECT_ROOT)} ({path.stat().st_size:,} bytes)")
        return 0
    problems = check()
    for problem in problems:
        print(f"MISMATCH {problem}")
    print("Fixtures match the real raw file layout." if not problems else
          "Regenerate with: python -m ingestion.ci_fixtures generate")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
