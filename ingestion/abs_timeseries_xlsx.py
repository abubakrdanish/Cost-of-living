"""Parse ABS 'Time Series Workbook' XLSX files into a long table.

Layout of each `Data*` sheet (verified on 6401.0 Table 9, Sep quarter 2025):
    row 0      : column A blank, then one series description per column
    rows 1-9   : metadata rows labelled in column A
                 (Unit, Series Type, Data Type, Frequency, Collection Month,
                  Series Start, Series End, No. Obs, Series ID)
    rows 10+   : column A is the period date, other columns are values
Blank cells mean "no observation" and are skipped.
"""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import openpyxl
import pandas as pd

METADATA_ROWS = {
    "Unit": "unit",
    "Series Type": "series_type",
    "Data Type": "data_type",
    "Frequency": "frequency",
    "Collection Month": "collection_month",
    "Series Start": "series_start",
    "Series End": "series_end",
    "No. Obs": "no_obs",
    "Series ID": "series_id",
}
FIRST_DATA_ROW = 1 + len(METADATA_ROWS)
COLUMNS = [
    "sheet", "series_id", "series_description", "unit", "series_type", "data_type",
    "frequency", "collection_month", "series_start", "series_end", "no_obs",
    "period", "value",
]


def looks_like_xlsx(content: bytes) -> bool:
    """XLSX files are ZIP archives; an HTML error page is not."""
    return content[:2] == b"PK"


def _as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _read_series_metadata(rows: list[tuple]) -> list[dict]:
    """One dict per data column (index >= 1) built from the header rows."""
    descriptions = rows[0]
    labels = [row[0] for row in rows[1:FIRST_DATA_ROW]]
    if labels != list(METADATA_ROWS):
        raise ValueError(f"Unexpected ABS metadata rows: {labels}")
    series = []
    for col in range(1, len(descriptions)):
        meta = {"series_description": (descriptions[col] or "").strip()}
        for offset, key in enumerate(METADATA_ROWS.values(), start=1):
            meta[key] = rows[offset][col]
        meta["series_start"] = _as_date(meta["series_start"])
        meta["series_end"] = _as_date(meta["series_end"])
        series.append(meta)
    return series


def parse_sheet(sheet_name: str, rows: list[tuple]) -> list[dict]:
    """Turn one Data sheet's rows into long records (one per non-blank cell)."""
    series = _read_series_metadata(rows)
    records = []
    for row in rows[FIRST_DATA_ROW:]:
        period = _as_date(row[0]) if row else None
        if period is None:
            continue  # trailing blank or footnote rows
        for col, meta in enumerate(series, start=1):
            value = row[col] if col < len(row) else None
            if value is None or value == "":
                continue
            records.append({"sheet": sheet_name, **meta, "period": period, "value": float(value)})
    return records


def parse_workbook(path: Path) -> pd.DataFrame:
    """Parse every `Data*` sheet in an ABS time series workbook."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        records: list[dict] = []
        for name in workbook.sheetnames:
            if name.startswith("Data"):
                rows = list(workbook[name].iter_rows(values_only=True))
                records.extend(parse_sheet(name, rows))
    finally:
        workbook.close()
    if not records:
        raise ValueError(f"No data found in {path.name}")
    return pd.DataFrame.from_records(records, columns=COLUMNS)
